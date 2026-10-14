from zipfile import ZipFile

import pytest
from docx import Document

from mane_worker.config import ROOT
from mane_worker.rendering import render_minutes


def test_docx_and_markdown_contain_evidence_and_actions(minutes, transcript, schema, tmp_path):
    docx, md = render_minutes(
        minutes,
        transcript,
        ROOT / "services/doc-worker/templates/minutes.docx",
        tmp_path,
        schema,
        "meeting-id",
    )
    document = Document(docx)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Pilot planning" in text
    assert "Evidence appendix" in text
    assert "We decided to launch" in text
    assert "{{TITLE}}" not in text
    assert len(document.tables) == 2
    assert "Prepare onboarding checklist" in document.tables[1].rows[1].cells[0].text
    markdown = md.read_text()
    assert "/app/meetings/meeting-id?t=0" in markdown
    assert "| Open |" in markdown
    with ZipFile(docx) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()


def test_render_rejects_invalid_evidence_after_review(minutes, transcript, schema, tmp_path):
    minutes["decisions"][0]["evidence"] = [{"start_s": 500, "end_s": 510}]
    with pytest.raises(ValueError, match="unsupported evidence"):
        render_minutes(
            minutes,
            transcript,
            ROOT / "services/doc-worker/templates/minutes.docx",
            tmp_path,
            schema,
        )
