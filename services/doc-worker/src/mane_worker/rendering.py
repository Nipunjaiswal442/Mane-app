from pathlib import Path

from docx import Document
from docx.shared import Inches

from mane_worker.models import Transcript
from mane_worker.pipeline import enforce_evidence, validate_minutes


def timestamp(seconds: float) -> str:
    return f"{int(seconds) // 60:02}:{int(seconds) % 60:02}"


def span_label(span: dict) -> str:
    return f"{timestamp(span['start_s'])}–{timestamp(span['end_s'])}"


def excerpts(span: dict, transcript: Transcript) -> str:
    return " ".join(
        segment.text
        for segment in transcript.segments
        if segment.end_s > span["start_s"] and segment.start_s < span["end_s"]
    )


def markdown_escape(value: str | None) -> str:
    return (
        (value or "—")
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\n", " ")
        .replace("\r", " ")
    )


def render_minutes(
    content: dict,
    transcript: Transcript,
    template: Path,
    directory: Path,
    schema: dict,
    meeting_id: str = "",
) -> tuple[Path, Path]:
    validate_minutes(content, schema)
    checked = enforce_evidence(content, transcript)
    # Approval must never silently publish unsupported claims added during review.
    if (
        checked["decisions"] != content["decisions"]
        or checked["action_items"] != content["action_items"]
    ):
        raise ValueError("Approved minutes contain unsupported evidence; return them to review")
    document = Document(template)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("{{TITLE}}", content["title"]).replace(
                "{{HELD_AT}}", content["held_at"] or "Date not specified"
            )
    directory.mkdir(parents=True, exist_ok=True)
    lines = [
        f"# {markdown_escape(content['title'])}",
        "",
        f"**Held:** {markdown_escape(content['held_at'])}",
        "",
        "## Attendees",
        "",
    ]
    document.add_heading("Attendees", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Name", "Identity"
    for attendee in content["attendees"]:
        cells = table.add_row().cells
        cells[0].text = attendee["name"]
        cells[1].text = "Matched organisation user" if attendee["matched_user_id"] else "Unverified"
        lines.append(f"- {markdown_escape(attendee['name'])}")
    if not content["attendees"]:
        lines.append("No attendees identified.")
    document.add_heading("Summary", level=1)
    document.add_paragraph(content["summary"])
    lines.extend(["", "## Summary", "", content["summary"], ""])
    document.add_heading("Agenda", level=1)
    lines.extend(["## Agenda", ""])
    for item in content["agenda"]:
        document.add_paragraph(item, style="List Bullet")
        lines.append(f"- {markdown_escape(item)}")
    document.add_heading("Decisions", level=1)
    lines.extend(["", "## Decisions", ""])
    for index, decision in enumerate(content["decisions"], 1):
        evidence = ", ".join(span_label(span) for span in decision["evidence"])
        document.add_paragraph(decision["text"], style="List Number")
        document.add_paragraph(f"Evidence: {evidence}", style="Caption")
        links = " · ".join(evidence_link(span, meeting_id) for span in decision["evidence"])
        lines.append(f"{index}. {markdown_escape(decision['text'])} — {links}")
    if not content["decisions"]:
        document.add_paragraph("No decisions recorded.")
        lines.append("No decisions recorded.")
    document.add_heading("Action items", level=1)
    action_table = document.add_table(rows=1, cols=5)
    action_table.style = "Table Grid"
    for cell, title in zip(
        action_table.rows[0].cells, ("Action", "Owner", "Due", "Status", "Evidence")
    ):
        cell.text = title
    lines.extend(
        [
            "",
            "## Action items",
            "",
            "| Action | Owner | Due | Status | Evidence |",
            "| --- | --- | --- | --- | --- |",
        ]
    )
    for item in content["action_items"]:
        values = [
            item["text"],
            item["owner"] or "Unassigned",
            item["due_date"] or "Not set",
            "Open",
            ", ".join(span_label(span) for span in item["evidence"]),
        ]
        for cell, value in zip(action_table.add_row().cells, values):
            cell.text = value
        values[-1] = " · ".join(evidence_link(span, meeting_id) for span in item["evidence"])
        lines.append("| " + " | ".join(markdown_escape(value) for value in values) + " |")
    for key, heading in (
        ("open_questions", "Open questions"),
        ("risks", "Risks"),
        ("confidence_notes", "Confidence notes"),
    ):
        document.add_heading(heading, level=1)
        lines.extend(["", f"## {heading}", ""])
        for text in content[key]:
            document.add_paragraph(text, style="List Bullet")
            lines.append(f"- {markdown_escape(text)}")
        if not content[key]:
            document.add_paragraph("None recorded.")
            lines.append("None recorded.")
    document.add_heading("Next meeting", level=1)
    document.add_paragraph(content["next_meeting"] or "Not scheduled.")
    lines.extend(["", "## Next meeting", "", content["next_meeting"] or "Not scheduled."])
    appendix = document.add_heading("Evidence appendix", level=1)
    appendix.paragraph_format.page_break_before = True
    lines.extend(["", "## Evidence appendix", ""])
    spans = {
        tuple((span["start_s"], span["end_s"]))
        for group in (content["decisions"], content["action_items"])
        for item in group
        for span in item["evidence"]
    }
    for start, end in sorted(spans):
        span = {"start_s": start, "end_s": end}
        document.add_heading(span_label(span), level=2)
        text = excerpts(span, transcript)
        document.add_paragraph(text)
        lines.extend([f"### {span_label(span)}", "", f"> {markdown_escape(text)}", ""])
    for section in document.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    docx_path, md_path = directory / "minutes.docx", directory / "minutes.md"
    document.save(docx_path)
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return docx_path, md_path


def evidence_link(span: dict, meeting_id: str) -> str:
    label = span_label(span)
    return (
        f"[{label}](/app/meetings/{meeting_id}?t={span['start_s']}&end={span['end_s']})"
        if meeting_id
        else label
    )
