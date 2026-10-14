"""Rebuild the checked-in template using python-docx; no network assets."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def create(destination: Path):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("111111")
    for name in ("Subtitle", "Caption"):
        document.styles[name].font.color.rgb = RGBColor.from_string("444444")
    document.styles["Title"].font.size = Pt(30)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(12)
    header = section.header.paragraphs[0]
    header.text = "MANE / MEETING RECORD"
    header.runs[0].font.size = Pt(9)
    header.runs[0].bold = True
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in {"val": "single", "sz": "12", "color": "111111", "space": "8"}.items():
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    header._p.get_or_add_pPr().append(border)
    document.add_paragraph("{{TITLE}}", style="Title")
    document.add_paragraph("{{HELD_AT}}", style="Subtitle")
    document.add_paragraph("Private · Reviewed meeting minutes", style="Caption")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("MANE · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "Mane"
    document.core_properties.subject = "Reviewed, evidence-grounded meeting minutes"
    document.save(destination)


if __name__ == "__main__":
    create(Path(__file__).with_name("minutes.docx"))
